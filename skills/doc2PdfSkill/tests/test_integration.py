from __future__ import annotations

import os
from pathlib import Path

import pytest

from doc2pdf_skill.converter import convert_document
from doc2pdf_skill.preflight import discover_libreoffice


pytestmark = pytest.mark.integration


@pytest.mark.skipif(os.environ.get("RUN_DOC2PDF_INTEGRATION") != "1", reason="set RUN_DOC2PDF_INTEGRATION=1")
@pytest.mark.parametrize("profile", ["fidelity", "balanced", "compact"])
def test_real_libreoffice_conversion(tmp_path: Path, profile: str) -> None:
    if discover_libreoffice() is None:
        pytest.skip("LibreOffice is unavailable")
    from docx import Document

    source = tmp_path / "book.docx"
    document = Document()
    document.add_heading("A conversion test", level=1)
    document.add_paragraph("Text with Romanian characters: ăâîșț.")
    document.add_table(rows=2, cols=2).cell(0, 0).text = "Table"
    document.save(source)
    result = convert_document(source, tmp_path / f"book-{profile}.pdf", profile=profile)
    assert result["status"] in {"passed", "passed_with_warnings"}
    assert Path(result["artifact"]).is_file()
    assert result["metrics"]["textCoverage"] >= 0.999
