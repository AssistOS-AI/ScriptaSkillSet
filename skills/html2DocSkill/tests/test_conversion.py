from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from bs4 import BeautifulSoup
from docx import Document
from lxml import etree
from PIL import Image
import pytest

from html2doc_skill.converter import convert_html, is_owned_docx
from html2doc_skill.models import Html2DocError
from html2doc_skill.validation import validate_docx


def _book(tmp_path: Path, *, notes: bool = False) -> Path:
    folder = tmp_path / "sample-book"
    assets = folder / "assets"
    assets.mkdir(parents=True)
    Image.new("RGB", (320, 180), "#336699").save(assets / "figure.png")
    (assets / "styles.css").write_text(
        """:root{--pdf-page-width:595.28pt;--pdf-page-height:841.89pt;--pdf-body-size:11pt}
        body{font-family:Georgia,serif;font-size:var(--pdf-body-size);color:#222}
        p{text-align:justify;margin-bottom:6pt} h1{font-size:2em}
        th{background-color:#ddeeff} img{width:240pt}""",
        encoding="utf-8",
    )
    note_markup = ""
    if notes:
        note_markup = """<p>A documented claim<a role="doc-noteref" href="#fn1">1</a> and an appendix<a role="doc-noteref" href="#en1">a</a>.</p>
        <aside id="fn1" role="doc-footnote">Footnote text.</aside>
        <aside id="en1" role="doc-endnote">Endnote text.</aside>"""
    source = folder / "index.html"
    source.write_text(
        f"""<!doctype html><html lang="en"><head><title>A Sample Book</title>
        <meta name="author" content="Ada Writer"><meta name="generator" content="pdf2html-skill">
        <link rel="stylesheet" href="assets/styles.css"></head><body>
        <main class="pdf-document" data-reader-content>
        <section class="source-page" id="page_1"><h1>A Sample Book</h1><p>A practical subtitle</p></section>
        <section class="source-page" id="page_2"><h2>Contents</h2><nav class="contents-list" aria-label="Contents"><p class="contents-list-entry">Chapter One 1</p></nav></section>
        <section class="source-page" id="page_3"><h1 id="chapter-one">Chapter One</h1>
        <p>Alpha <strong>bold</strong>, <em>italic</em>, and <a href="https://example.com">linked</a> text.</p>
        <ul><li>First item<ul><li>Nested item</li></ul></li></ul>{note_markup}</section>
        <section class="source-page" id="page_4"><h1 id="chapter-two" style="line-height:1.5">Chapter Two</h1>
        <p>Second chapter links to <a href="#chapter-one">the first chapter</a>.</p>
        <a href="https://example.org/redundant"><p><a href="https://example.org/redundant">Nested source link</a></p></a>
        <figure><img src="assets/figure.png" alt="Blue sample"><figcaption>Figure 1. Sample.</figcaption></figure>
        <table><thead><tr><th colspan="2">Measurements</th></tr></thead><tbody><tr><td>A</td><td>42</td></tr></tbody></table>
        </section></main><script>window.reader=true</script></body></html>""",
        encoding="utf-8",
    )
    return source


def test_converts_scripta_book_to_native_docx(tmp_path: Path) -> None:
    source = _book(tmp_path)
    result = convert_html(source)
    artifact = Path(result["artifact"])
    assert artifact == source.with_name("sample-book.docx")
    assert result["status"] == "passed"
    assert result["metrics"]["textCoverage"] == 1.0
    assert result["metrics"]["images"] == 1
    assert result["metrics"]["tables"] == 1
    assert is_owned_docx(artifact)

    document = Document(artifact)
    assert document.core_properties.title == "A Sample Book"
    assert document.core_properties.author == "Ada Writer"
    assert len(document.sections) == 2
    assert document.sections[0].different_first_page_header_footer is False
    assert "Chapter One" in " ".join(paragraph.text for paragraph in document.paragraphs)
    assert next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text == "Chapter Two" and paragraph.style.name == "Heading 1"
    ).paragraph_format.line_spacing == 1.15
    assert next(paragraph for paragraph in document.paragraphs if paragraph.text == "Contents").style.name == "TOC Heading"
    with ZipFile(artifact) as package:
        xml = package.read("word/document.xml").decode("utf-8")
        settings = package.read("word/settings.xml").decode("utf-8")
        headers = [name for name in package.namelist() if name.startswith("word/header")]
        footer_names = [name for name in package.namelist() if name.startswith("word/footer")]
        footers = "".join(package.read(name).decode("utf-8") for name in footer_names)
        assert "Update field to build the table of contents" not in xml
        assert headers == []
        assert len(footer_names) == 1
        assert "PAGE" in footers
        assert xml.count('w:fmt="decimal" w:start="1"') == 1
        assert xml.count("ScriptaTOC") == 2
        assert "w:updateFields" in settings
        assert "w:bookmarkStart" in xml and "hyperlink" in xml


def test_native_footnotes_and_endnotes_have_complete_relationships(tmp_path: Path) -> None:
    source = _book(tmp_path, notes=True)
    result = convert_html(source)
    artifact = Path(result["artifact"])
    assert result["metrics"]["footnotes"] == 1
    assert result["metrics"]["endnotes"] == 1
    with ZipFile(artifact) as package:
        assert "word/footnotes.xml" in package.namelist()
        assert "word/endnotes.xml" in package.namelist()
        assert b"Footnote text." in package.read("word/footnotes.xml")
        assert b"Endnote text." in package.read("word/endnotes.xml")
        relationships = package.read("word/_rels/document.xml.rels")
        assert b"/footnotes" in relationships and b"/endnotes" in relationships
    assert validate_docx(source, artifact)["status"] == "passed"


def test_rejects_non_scripta_html_and_remote_assets(tmp_path: Path) -> None:
    ordinary = tmp_path / "ordinary.html"
    ordinary.write_text("<html><body><h1>Not Scripta</h1></body></html>", encoding="utf-8")
    with pytest.raises(Html2DocError, match=r"main\[data-reader-content\]"):
        convert_html(ordinary)

    remote = tmp_path / "remote.html"
    remote.write_text("<html><body><main data-reader-content><img src='https://example.com/x.png'></main></body></html>", encoding="utf-8")
    with pytest.raises(Html2DocError, match="Remote resources"):
        convert_html(remote)


def test_overwrite_only_replaces_owned_output(tmp_path: Path) -> None:
    source = _book(tmp_path)
    artifact = Path(convert_html(source)["artifact"])
    with pytest.raises(Html2DocError, match="Use --overwrite"):
        convert_html(source)
    assert Path(convert_html(source, overwrite=True)["artifact"]) == artifact

    unrelated = tmp_path / "unrelated.docx"
    Document().save(unrelated)
    with pytest.raises(Html2DocError, match="not owned"):
        convert_html(source, unrelated, overwrite=True)


def test_validation_detects_content_loss(tmp_path: Path) -> None:
    source = _book(tmp_path)
    artifact = Path(convert_html(source)["artifact"])
    document = Document(artifact)
    for paragraph in document.paragraphs:
        if "Second chapter" in paragraph.text:
            paragraph.text = "Removed"
    damaged = tmp_path / "damaged.docx"
    document.save(damaged)
    report = validate_docx(source, damaged)
    assert report["status"] == "failed"
    assert any(item["code"] in {"text-coverage", "ownership-marker"} for item in report["findings"])


def test_metadata_overrides_and_toc_insertion_without_source_toc(tmp_path: Path) -> None:
    source = _book(tmp_path)
    soup = BeautifulSoup(source.read_text(encoding="utf-8"), "html.parser")
    soup.select_one("section#page_2").decompose()
    source.write_text(str(soup), encoding="utf-8")
    result = convert_html(source, title="Override Title", author="New Author", language="ro")
    document = Document(result["artifact"])
    assert document.core_properties.title == "Override Title"
    assert document.core_properties.author == "New Author"
    assert document.core_properties.language == "ro"
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts.index("A Sample Book") < texts.index("Contents")
    assert texts.index("Contents") < texts.index("Chapter One")
    with ZipFile(result["artifact"]) as package:
        xml = package.read("word/document.xml")
        assert b"ScriptaTOC" in xml
        assert b"Update field to build the table of contents" not in xml


def test_document_without_headings_still_places_toc_before_content(tmp_path: Path) -> None:
    source = tmp_path / "plain.html"
    source.write_text("<html><body><main data-reader-content><p>Plain document body.</p></main></body></html>", encoding="utf-8")
    result = convert_html(source)
    document = Document(result["artifact"])
    assert document.paragraphs[0].text == "Contents"
    assert document.paragraphs[-1].text == "Plain document body."
    assert result["status"] == "passed_with_warnings"


def test_plain_scripta_contents_page_is_replaced_and_slug_title_is_not_used(tmp_path: Path) -> None:
    source = _book(tmp_path)
    soup = BeautifulSoup(source.read_text(encoding="utf-8"), "html.parser")
    soup.title.string = "A_Sample_Book"
    nav = soup.select_one("nav.contents-list")
    page = nav.find_parent("section")
    page["aria-label"] = "PDF page 2"
    first_entry = nav.find("p")
    first_entry.clear()
    strong = soup.new_tag("strong")
    strong.string = "Chapter One"
    first_entry.append(strong)
    second_entry = soup.new_tag("p")
    second_entry.string = "Chapter Two"
    nav.append(second_entry)
    chapter_one = soup.select_one("section#page_3 h1")
    subsection = soup.new_tag("h2")
    subsection.string = "An unlisted subsection"
    chapter_one.insert_after(subsection)
    nav.unwrap()
    source.write_text(str(soup), encoding="utf-8")
    result = convert_html(source)
    document = Document(result["artifact"])
    assert document.core_properties.title == "A Sample Book"
    assert [paragraph.text for paragraph in document.paragraphs].count("Contents") == 1
    assert "PDF page 2" not in [paragraph.text for paragraph in document.paragraphs]
    with ZipFile(result["artifact"]) as package:
        headers = [name for name in package.namelist() if name.startswith("word/header")]
        xml = package.read("word/document.xml").decode("utf-8")
        root = etree.fromstring(xml.encode("utf-8"))
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        toc_paragraphs = root.xpath(".//w:p[w:pPr/w:pStyle[starts-with(@w:val, 'ScriptaTOC')]]", namespaces=namespace)
        assert headers == []
        assert "Update field to build the table of contents" not in xml
        toc_emphasis = {
            "".join(paragraph.xpath(".//w:t/text()", namespaces=namespace)): bool(paragraph.xpath(".//w:b", namespaces=namespace))
            for paragraph in toc_paragraphs
        }
        assert toc_emphasis["Chapter One"] is True
        assert toc_emphasis["An unlisted subsection"] is False
        assert toc_emphasis["Chapter Two"] is False


def test_front_matter_source_pages_remain_separate(tmp_path: Path) -> None:
    source = tmp_path / "front-matter.html"
    source.write_text(
        """<html><body><main data-reader-content>
        <section class="source-page"><p>Copyright page</p></section>
        <section class="source-page"><h1>Book Title</h1><p>Book subtitle</p></section>
        <section class="source-page"><h1>A Note</h1><p>Note body.</p></section>
        <section class="source-page" aria-label="PDF page 4"><h1>Contents</h1><p>Prologue</p></section>
        <section class="source-page"><h1>Prologue</h1><p>Body text.</p></section>
        </main></body></html>""",
        encoding="utf-8",
    )
    result = convert_html(source)
    document = Document(result["artifact"])
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    assert paragraphs["Book Title"].style.name == "Scripta Book Title"
    assert paragraphs["Book subtitle"].style.name == "Scripta Book Subtitle"
    assert paragraphs["Book Title"].paragraph_format.page_break_before is True
    assert paragraphs["A Note"].paragraph_format.page_break_before is True
    assert paragraphs["Contents"].paragraph_format.page_break_before is True
    assert "PDF page 4" not in paragraphs
