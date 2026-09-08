from __future__ import annotations

from pathlib import Path
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.text.paragraph import Paragraph
from lxml import etree


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": WORD_NS, "r": REL_NS}
FOOTNOTE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
ENDNOTE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes"


def field_run(paragraph: Paragraph, instruction: str, placeholder: str = "") -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(instruction_node)
    paragraph.add_run()._r.append(separate)
    paragraph.add_run(placeholder)
    paragraph.add_run()._r.append(end)


def add_hyperlink(paragraph: Paragraph, text: str, href: str, *, internal: bool = False, bold: bool = False, italic: bool = False) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if internal:
        hyperlink.set(qn("w:anchor"), href)
    else:
        relation = paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    if bold:
        properties.append(OxmlElement("w:b"))
    if italic:
        properties.append(OxmlElement("w:i"))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def bookmark_name(name: str) -> str:
    return "_" + "".join(character if character.isalnum() or character == "_" else "_" for character in name)[:38]


def add_bookmark(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    safe = bookmark_name(name)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), safe)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def page_number_format(section, fmt: str, start: int = 1) -> None:
    section_properties = section._sectPr
    existing = section_properties.find(qn("w:pgNumType"))
    if existing is None:
        existing = OxmlElement("w:pgNumType")
        section_properties.append(existing)
    existing.set(qn("w:fmt"), fmt)
    existing.set(qn("w:start"), str(start))


def enable_update_fields(document) -> None:
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    properties.append(node)


def _note_part(kind: str, notes: dict[int, str]) -> bytes:
    root = etree.Element(f"{{{WORD_NS}}}{kind}s", nsmap={"w": WORD_NS})
    for note_id, note_type, marker in ((-1, "separator", "separator"), (-2, "continuationSeparator", "continuationSeparator")):
        note = etree.SubElement(root, f"{{{WORD_NS}}}{kind}", {f"{{{WORD_NS}}}id": str(note_id), f"{{{WORD_NS}}}type": note_type})
        paragraph = etree.SubElement(note, f"{{{WORD_NS}}}p")
        run = etree.SubElement(paragraph, f"{{{WORD_NS}}}r")
        etree.SubElement(run, f"{{{WORD_NS}}}{marker}")
    for note_id, text in notes.items():
        note = etree.SubElement(root, f"{{{WORD_NS}}}{kind}", {f"{{{WORD_NS}}}id": str(note_id)})
        paragraph = etree.SubElement(note, f"{{{WORD_NS}}}p")
        ppr = etree.SubElement(paragraph, f"{{{WORD_NS}}}pPr")
        etree.SubElement(ppr, f"{{{WORD_NS}}}pStyle", {f"{{{WORD_NS}}}val": "FootnoteText" if kind == "footnote" else "EndnoteText"})
        run_ref = etree.SubElement(paragraph, f"{{{WORD_NS}}}r")
        rpr = etree.SubElement(run_ref, f"{{{WORD_NS}}}rPr")
        etree.SubElement(rpr, f"{{{WORD_NS}}}rStyle", {f"{{{WORD_NS}}}val": "FootnoteReference" if kind == "footnote" else "EndnoteReference"})
        etree.SubElement(run_ref, f"{{{WORD_NS}}}{kind}Ref")
        run = etree.SubElement(paragraph, f"{{{WORD_NS}}}r")
        text_node = etree.SubElement(run, f"{{{WORD_NS}}}t", {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
        text_node.text = f" {text}"
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def patch_notes(path: Path, footnotes: dict[int, str], endnotes: dict[int, str]) -> None:
    if not footnotes and not endnotes:
        return
    with ZipFile(path) as source:
        files = {name: source.read(name) for name in source.namelist()}
    document = etree.fromstring(files["word/document.xml"])
    for kind, notes in (("footnote", footnotes), ("endnote", endnotes)):
        if not notes:
            continue
        for run in list(document.xpath(".//w:r", namespaces=NS)):
            texts = run.xpath("./w:t/text()", namespaces=NS)
            if len(texts) != 1 or not texts[0].startswith(f"HTML2DOC_{kind.upper()}_"):
                continue
            note_id = int(texts[0].rsplit("_", 1)[-1])
            replacement = etree.Element(f"{{{WORD_NS}}}r")
            rpr = etree.SubElement(replacement, f"{{{WORD_NS}}}rPr")
            etree.SubElement(rpr, f"{{{WORD_NS}}}rStyle", {f"{{{WORD_NS}}}val": "FootnoteReference" if kind == "footnote" else "EndnoteReference"})
            etree.SubElement(replacement, f"{{{WORD_NS}}}{kind}Reference", {f"{{{WORD_NS}}}id": str(note_id)})
            run.getparent().replace(run, replacement)
        files[f"word/{kind}s.xml"] = _note_part(kind, notes)

        relationships = etree.fromstring(files["word/_rels/document.xml.rels"])
        used = [int(value[3:]) for value in relationships.xpath("./pr:Relationship/@Id", namespaces={"pr": PACKAGE_REL_NS}) if value.startswith("rId") and value[3:].isdigit()]
        relationship = etree.SubElement(relationships, f"{{{PACKAGE_REL_NS}}}Relationship")
        relationship.set("Id", f"rId{max(used, default=0) + 1}")
        relationship.set("Type", FOOTNOTE_REL if kind == "footnote" else ENDNOTE_REL)
        relationship.set("Target", f"{kind}s.xml")
        files["word/_rels/document.xml.rels"] = etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone="yes")

        content_types = etree.fromstring(files["[Content_Types].xml"])
        override = etree.SubElement(content_types, f"{{{CONTENT_NS}}}Override")
        override.set("PartName", f"/word/{kind}s.xml")
        override.set("ContentType", f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml")
        files["[Content_Types].xml"] = etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone="yes")

    files["word/document.xml"] = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone="yes")
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(temporary, "w", ZIP_DEFLATED) as target:
            for name, payload in files.items():
                target.writestr(name, payload)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
