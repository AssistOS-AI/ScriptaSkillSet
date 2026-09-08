from __future__ import annotations

from collections import Counter
from pathlib import Path
import re
import tempfile

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from .css import StyleResolver, color_hex, local_asset, points
from .models import BuildState, DocumentMetadata, Finding, Html2DocError
from .ooxml import (
    add_bookmark,
    add_hyperlink,
    bookmark_name,
    enable_update_fields,
    field_run,
    page_number_format,
    patch_notes,
    set_cell_shading,
    set_repeat_table_header,
)


WORD_RE = re.compile(r"\w+", re.UNICODE)
BLOCKS = {"address", "blockquote", "dd", "div", "dl", "dt", "figcaption", "figure", "h1", "h2", "h3", "h4", "h5", "h6", "hr", "li", "nav", "ol", "p", "pre", "section", "table", "ul"}
SKIP = {"script", "style", "template", "noscript"}


def _tokens(text: str) -> list[str]:
    return WORD_RE.findall(text.casefold().replace("\u00ad", ""))


def _text_counts(source: list[str], output: list[str]) -> tuple[float, float]:
    if not source:
        return 1.0, 1.0
    source_counts, output_counts = Counter(source), Counter(output)
    coverage = sum(min(count, output_counts[token]) for token, count in source_counts.items()) / len(source)
    if len(source) < 2:
        return coverage, 1.0
    source_pairs, output_pairs = Counter(zip(source, source[1:])), Counter(zip(output, output[1:]))
    order = sum(min(count, output_pairs[pair]) for pair, count in source_pairs.items()) / sum(source_pairs.values())
    return coverage, order


def _is_hidden(tag: Tag) -> bool:
    if tag.name in SKIP or tag.has_attr("hidden") or str(tag.get("aria-hidden", "")).casefold() == "true":
        return True
    style = str(tag.get("style", "")).replace(" ", "").casefold()
    return "display:none" in style or "visibility:hidden" in style


def find_toc_container(soup: BeautifulSoup) -> Tag | None:
    structured = soup.select_one("table.toc-table, nav.contents-list")
    if structured is not None:
        return structured
    for heading in soup.find_all(["h1", "h2", "h3"]):
        if heading.get_text(" ", strip=True).casefold() not in {"contents", "table of contents"}:
            continue
        page = heading.find_parent("section", class_="source-page")
        if page is not None:
            return page
        return heading
    return None


def find_cover_page(soup: BeautifulSoup) -> Tag | None:
    first = soup.select_one("section.source-page")
    if not first:
        return None
    if "source-page-full-image" in first.get("class", []):
        return first
    blocks = first.find_all(["h1", "h2", "p", "figure"], recursive=False)
    return first if first.find("h1") and len(blocks) <= 5 else None


def find_body_heading(soup: BeautifulSoup, toc: Tag | None, cover: Tag | None) -> Tag | None:
    headings = soup.find_all(["h1", "h2"])
    if toc:
        return next(
            (heading for heading in toc.find_all_next(["h1", "h2"]) if toc not in heading.parents),
            None,
        )
    if cover:
        return next((heading for heading in headings if cover not in heading.parents), None)
    return headings[0] if headings else None


def toc_entry_headings(soup: BeautifulSoup, body_heading: Tag | None) -> list[Tag]:
    if body_heading is None:
        return []
    headings = soup.find_all(["h1", "h2", "h3"])
    body_id = id(body_heading)
    start = next((index for index, heading in enumerate(headings) if id(heading) == body_id), len(headings))
    return headings[start:]


def toc_label(toc: Tag | None) -> str:
    if toc is None:
        return "Contents"
    heading = toc if toc.name in {"h1", "h2", "h3"} else toc.find(["h1", "h2", "h3"])
    if heading is not None and heading.get_text(" ", strip=True).casefold() in {"contents", "table of contents"}:
        return heading.get_text(" ", strip=True)
    if toc.name in {"nav", "table"}:
        label = str(toc.get("aria-label", "")).strip()
        if label and not re.fullmatch(r"PDF page \d+", label, flags=re.I):
            return label
    return "Contents"


def toc_entry_boldness(toc: Tag | None, headings: list[Tag]) -> list[bool]:
    """Match source TOC lines to document headings and retain their emphasis."""
    if toc is None:
        return [int(heading.name[1]) == 1 for heading in headings]
    candidates = toc.find_all(["p", "li"], recursive=True)
    if not candidates:
        return [int(heading.name[1]) == 1 for heading in headings]
    result: list[bool] = []
    cursor = 0
    for heading in headings:
        wanted = _tokens(heading.get_text(" ", strip=True))
        collected: list[str] = []
        bold_tokens = 0
        total_tokens = 0
        matched = False
        candidate_cursor = cursor
        while candidate_cursor < len(candidates):
            candidate = candidates[candidate_cursor]
            candidate_cursor += 1
            tokens = _tokens(candidate.get_text(" ", strip=True))
            if not tokens:
                continue
            proposed = collected + tokens
            if wanted[:len(proposed)] != proposed:
                if not collected:
                    continue
                break
            collected = proposed
            total_tokens += len(tokens)
            bold_tokens += sum(len(_tokens(strong.get_text(" ", strip=True))) for strong in candidate.find_all(["strong", "b"]))
            if collected == wanted:
                matched = True
                break
        if matched:
            cursor = candidate_cursor
        result.append((bold_tokens / total_tokens >= 0.6) if matched and total_tokens else int(heading.name[1]) == 1)
    return result


def semantic_text(soup: BeautifulSoup) -> str:
    clone = BeautifulSoup(str(soup), "html.parser")
    root = clone.select_one("main[data-reader-content]") or clone.body or clone
    footnotes = [tag.get_text(" ", strip=True) for tag in root.select('[role="doc-footnote"]')]
    endnotes = [tag.get_text(" ", strip=True) for tag in root.select('[role="doc-endnote"]')]
    for tag in list(clone.find_all(True)):
        if _is_hidden(tag):
            tag.decompose()
    toc = find_toc_container(clone)
    if toc is not None:
        toc.decompose()
    for tag in clone.select("a[role=doc-noteref], [role=doc-footnote], [role=doc-endnote]"):
        tag.decompose()
    return " ".join([root.get_text(" ", strip=True), *footnotes, *endnotes]).strip()


def _metadata(soup: BeautifulSoup, source: Path, *, title: str | None, author: str | None, language: str | None, resolver: StyleResolver) -> DocumentMetadata:
    title_tag = soup.find("title")
    first_h1 = soup.find("h1")
    resolved_title = title or (first_h1.get_text(" ", strip=True) if first_h1 else "") or (title_tag.get_text(" ", strip=True) if title_tag else "") or source.parent.name
    author_tag = soup.find("meta", attrs={"name": re.compile(r"^author$", re.I)})
    resolved_author = author if author is not None else (str(author_tag.get("content", "")) if author_tag else "")
    resolved_language = language or str(soup.html.get("lang", "") if soup.html else "") or "und"
    root_style = resolver.style(soup.html or soup)
    width = points(root_style.get("--pdf-page-width"), page=612) or 595.28
    height = points(root_style.get("--pdf-page-height"), page=792) or 841.89
    body = soup.body or soup
    body_size = points(resolver.resolve(body, "font-size"), base=11) or points(root_style.get("--pdf-body-size"), base=11) or 11
    family = (resolver.resolve(body, "font-family") or "Georgia").split(",")[0].strip(" '\"")
    if family.casefold() in {"serif", "sans-serif", "system-ui"}:
        family = "Georgia" if family.casefold() == "serif" else "Arial"
    page_styles = [resolver.style(tag) for tag in soup.select("section.source-page")]
    margins = []
    for style in page_styles:
        for name in ("--pdf-page-left", "--pdf-page-right", "--pdf-page-top"):
            value = points(style.get(name), page=width)
            if value and 18 <= value <= width * 0.25:
                margins.append(value)
    margin = sorted(margins)[len(margins) // 2] if margins else 72.0
    return DocumentMetadata(resolved_title, resolved_author, resolved_language, width, height, margin, family, max(8, min(14, body_size)))


def _set_font(style, name: str, size: float, *, bold: bool | None = None, italic: bool | None = None) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def _styles(document: Document, metadata: DocumentMetadata) -> None:
    styles = document.styles
    normal = styles["Normal"]
    _set_font(normal, metadata.body_font, metadata.body_size_pt)
    normal.paragraph_format.space_after = Pt(metadata.body_size_pt * 0.55)
    normal.paragraph_format.line_spacing = 1.25
    for level, size in enumerate((24, 18, 15, 13, 12, 11), start=1):
        style = styles[f"Heading {level}"]
        _set_font(style, metadata.body_font, size, bold=True)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(18 if level <= 2 else 12)
        style.paragraph_format.space_after = Pt(6)
    definitions = (
        ("Scripta Book Title", 30, True, False), ("Scripta Book Subtitle", 16, False, True),
        ("Code Block", 9, False, False), ("Book Quote", metadata.body_size_pt, False, True),
        ("Scripta TOC 1", metadata.body_size_pt, False, False),
        ("Scripta TOC 2", metadata.body_size_pt, False, False),
        ("Scripta TOC 3", max(8, metadata.body_size_pt - 0.5), False, False),
    )
    for name, size, bold, italic in definitions:
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_font(style, "Consolas" if name == "Code Block" else metadata.body_font, size, bold=bold, italic=italic)
    styles["Scripta Book Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Scripta Book Title"].paragraph_format.space_before = Pt(120)
    styles["Scripta Book Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Book Quote"].paragraph_format.left_indent = Inches(0.4)
    styles["Book Quote"].paragraph_format.right_indent = Inches(0.4)
    for level in range(1, 4):
        toc_style = styles[f"Scripta TOC {level}"]
        toc_style.paragraph_format.left_indent = Inches(0.25 * (level - 1))
        toc_style.paragraph_format.space_after = Pt(5 if level == 1 else 3)
    styles["Caption"].font.name = metadata.body_font
    styles["Caption"].font.size = Pt(max(8, metadata.body_size_pt - 1.5))


def _page_setup(section, metadata: DocumentMetadata) -> None:
    section.page_width = Pt(metadata.page_width_pt)
    section.page_height = Pt(metadata.page_height_pt)
    if metadata.page_width_pt > metadata.page_height_pt:
        section.orientation = WD_ORIENT.LANDSCAPE
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(metadata.margin_pt)
    section.header_distance = Pt(28)
    section.footer_distance = Pt(28)


def _configure_front(section) -> None:
    page_number_format(section, "lowerRoman", 1)
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field_run(paragraph, "PAGE", "1")


def _configure_body(section) -> None:
    page_number_format(section, "decimal", 1)


class Builder:
    def __init__(self, source: Path, soup: BeautifulSoup, resolver: StyleResolver, metadata: DocumentMetadata):
        self.source, self.soup, self.resolver, self.metadata = source, soup, resolver, metadata
        self.document = Document()
        self.state = BuildState(source, warnings=list(resolver.warnings))
        self.pending_bookmarks: list[str] = []
        _styles(self.document, metadata)
        _page_setup(self.document.sections[0], metadata)
        _configure_front(self.document.sections[0])
        self.document.core_properties.title = metadata.title
        self.document.core_properties.author = metadata.author
        self.document.core_properties.language = metadata.language
        self.document.core_properties.keywords = "html2doc-skill; Scripta"
        enable_update_fields(self.document)
        self.note_targets: dict[str, tuple[str, str]] = {}
        for tag in soup.find_all(True):
            role = str(tag.get("role", "")).casefold()
            if tag.get("id") and role in {"doc-footnote", "doc-endnote"}:
                self.note_targets[str(tag["id"])] = ("footnote" if role == "doc-footnote" else "endnote", tag.get_text(" ", strip=True))
        for reference in soup.select('a[role="doc-noteref"][href^="#"]'):
            target = str(reference.get("href", ""))[1:]
            if target not in self.note_targets:
                raise Html2DocError(f"Semantic note reference targets a missing footnote/endnote: #{target}")
        toc = find_toc_container(soup)
        self.toc_tag_id = id(toc) if toc else None
        toc_page = toc if toc is not None and "source-page" in toc.get("class", []) else (toc.find_parent("section", class_="source-page") if toc else None)
        self.toc_page_id = id(toc_page) if toc_page else None
        self.cover_page = find_cover_page(soup)
        title_heading = soup.find("h1")
        self.title_heading_id = id(title_heading) if title_heading else None
        self.title_page = title_heading.find_parent("section", class_="source-page") if title_heading else None
        candidate = find_body_heading(soup, toc, self.cover_page)
        self.body_heading_id = id(candidate) if candidate else None
        body_page = candidate.find_parent("section", class_="source-page") if candidate else None
        self.body_page_id = id(body_page) if body_page else None
        self.toc_entries = toc_entry_headings(soup, candidate)
        self.toc_boldness = toc_entry_boldness(toc, self.toc_entries)
        self.front_source_pages_emitted = 0
        self.page_break_pending = False
        for index, heading in enumerate(self.toc_entries, start=1):
            if not heading.get("id"):
                heading["id"] = f"html2doc-heading-{index:04d}"

    def insert_toc(self, label: str = "Contents") -> None:
        if self.state.toc_inserted:
            return
        previous = self.document.paragraphs[-1] if self.document.paragraphs else None
        has_heading = previous is not None and previous.text.strip().casefold() in {"contents", "table of contents", label.strip().casefold()}
        if not has_heading:
            heading = self.document.add_paragraph(label, style="TOC Heading")
            heading.paragraph_format.page_break_before = bool(self.document.paragraphs[:-1])
        else:
            previous.style = "TOC Heading"
        self.page_break_pending = False
        for heading, is_bold in zip(self.toc_entries, self.toc_boldness):
            level = min(3, int(heading.name[1]))
            paragraph = self.document.add_paragraph(style=f"Scripta TOC {level}")
            add_hyperlink(
                paragraph,
                heading.get_text(" ", strip=True),
                bookmark_name(str(heading["id"])),
                internal=True,
                bold=is_bold,
            )
            self.state.hyperlink_count += 1
        self.state.toc_inserted = True

    def start_body(self) -> None:
        if self.state.body_section_created:
            return
        if not self.state.toc_inserted:
            self.insert_toc()
        section = self.document.add_section(WD_SECTION.NEW_PAGE)
        _page_setup(section, self.metadata)
        _configure_body(section)
        self.state.body_section_created = True

    def add_inline(self, paragraph, node, *, bold: bool = False, italic: bool = False, code: bool = False) -> None:
        if isinstance(node, NavigableString):
            text = str(node) if code else re.sub(r"\s+", " ", str(node))
            if text:
                run = paragraph.add_run(text)
                run.bold, run.italic = bold, italic
                if code:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
            return
        if not isinstance(node, Tag) or _is_hidden(node):
            return
        name = node.name.casefold()
        if name == "br":
            paragraph.add_run().add_break()
            return
        if name == "img":
            self.add_image(node, paragraph)
            return
        if name == "a":
            href = str(node.get("href", ""))
            role = str(node.get("role", "")).casefold()
            target = href[1:] if href.startswith("#") else ""
            if role == "doc-noteref" and target in self.note_targets:
                kind, text = self.note_targets[target]
                notes = self.state.footnotes if kind == "footnote" else self.state.endnotes
                note_id = len(notes) + 1
                notes[note_id] = text
                paragraph.add_run(f"HTML2DOC_{kind.upper()}_{note_id}")
                return
            label = node.get_text(" ", strip=False)
            if href and label:
                add_hyperlink(paragraph, label, bookmark_name(target) if target else href, internal=bool(target), bold=bold, italic=italic)
                self.state.hyperlink_count += 1
            else:
                paragraph.add_run(label)
            return
        next_bold = bold or name in {"b", "strong"}
        next_italic = italic or name in {"em", "i", "cite"}
        for child in node.children:
            before = len(paragraph.runs)
            self.add_inline(paragraph, child, bold=next_bold, italic=next_italic, code=code or name == "code")
            if len(paragraph.runs) > before:
                run = paragraph.runs[-1]
                if name == "u":
                    run.underline = True
                elif name in {"s", "del"}:
                    run.font.strike = True
                elif name == "sup":
                    run.font.superscript = True
                elif name == "sub":
                    run.font.subscript = True

    def apply_paragraph_style(self, paragraph, tag: Tag) -> None:
        style = self.resolver.style(tag)
        alignment = self.resolver.resolve(tag, "text-align")
        paragraph.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT, "start": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "end": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }.get(str(alignment).casefold(), paragraph.alignment)
        base = self.metadata.body_size_pt
        for css_name, attr in (("margin-top", "space_before"), ("margin-bottom", "space_after"), ("text-indent", "first_line_indent"), ("margin-left", "left_indent"), ("margin-right", "right_indent")):
            value = points(self.resolver.resolve(tag, css_name), base=base, page=self.metadata.page_width_pt)
            if value is not None and abs(value) <= self.metadata.page_width_pt / 2:
                setattr(paragraph.paragraph_format, attr, Pt(value))
        line_height = self.resolver.resolve(tag, "line-height")
        if line_height and re.fullmatch(r"[0-9.]+", line_height):
            maximum = 1.15 if tag.name in {"h1", "h2", "h3", "h4", "h5", "h6"} else 3
            paragraph.paragraph_format.line_spacing = max(0.8, min(maximum, float(line_height)))
        if str(self.resolver.resolve(tag, "break-before") or self.resolver.resolve(tag, "page-break-before")).casefold() in {"page", "always"}:
            paragraph.paragraph_format.page_break_before = True
        color = color_hex(self.resolver.resolve(tag, "color"))
        size = points(self.resolver.resolve(tag, "font-size"), base=base)
        family = self.resolver.resolve(tag, "font-family")
        weight = str(self.resolver.resolve(tag, "font-weight") or "").casefold()
        font_style = str(self.resolver.resolve(tag, "font-style") or "").casefold()
        for run in paragraph.runs:
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            if size:
                run.font.size = Pt(max(6, min(72, size)))
            if family:
                run.font.name = family.split(",")[0].strip(" '\"")
            if weight in {"bold", "bolder", "600", "700", "800", "900"}:
                run.bold = True
            if font_style in {"italic", "oblique"}:
                run.italic = True

    def add_pending_bookmarks(self, paragraph) -> None:
        while self.pending_bookmarks:
            name = self.pending_bookmarks.pop(0)
            if name in self.state.bookmark_ids:
                continue
            bookmark_id = len(self.state.bookmark_ids) + 1
            self.state.bookmark_ids[name] = bookmark_id
            add_bookmark(paragraph, name, bookmark_id)

    def apply_pending_page_break(self, paragraph) -> None:
        if self.page_break_pending:
            paragraph.paragraph_format.page_break_before = True
            self.page_break_pending = False

    def add_paragraph(self, tag: Tag, style: str | None = None):
        paragraph = self.document.add_paragraph(style=style)
        self.apply_pending_page_break(paragraph)
        self.add_pending_bookmarks(paragraph)
        if tag.get("id"):
            bookmark_id = len(self.state.bookmark_ids) + 1
            self.state.bookmark_ids[str(tag["id"])] = bookmark_id
            add_bookmark(paragraph, str(tag["id"]), bookmark_id)
        for child in tag.children:
            self.add_inline(paragraph, child)
        self.apply_paragraph_style(paragraph, tag)
        return paragraph

    def add_image(self, tag: Tag, paragraph=None) -> None:
        source = str(tag.get("src", ""))
        if not source:
            raise Html2DocError("An image has no src attribute.")
        path = local_asset(self.source, source)
        intrinsic_width = available = self.metadata.page_width_pt - 2 * self.metadata.margin_pt
        try:
            with Image.open(path) as image:
                dpi_x = float(image.info.get("dpi", (96, 96))[0] or 96)
                intrinsic_width = image.width * 72 / dpi_x
                image.verify()
        except Exception as error:
            raise Html2DocError(f"Invalid or unsupported image {source}: {error}") from error
        target = paragraph or self.document.add_paragraph()
        if paragraph is None:
            self.apply_pending_page_break(target)
        self.add_pending_bookmarks(target)
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width = points(self.resolver.resolve(tag, "width"), base=self.metadata.body_size_pt, page=available)
        width = min(available, width or intrinsic_width)
        run = target.add_run()
        shape = run.add_picture(str(path), width=Pt(max(36, width)))
        alt = str(tag.get("alt", ""))
        if alt:
            shape._inline.docPr.set("descr", alt)
        self.state.image_count += 1

    def add_table(self, tag: Tag) -> None:
        caption = tag.find("caption", recursive=False)
        if caption:
            self.add_paragraph(caption, "Caption")
        if self.pending_bookmarks or self.page_break_pending:
            marker = self.document.add_paragraph()
            marker.paragraph_format.space_after = Pt(0)
            marker.paragraph_format.line_spacing = 1
            self.apply_pending_page_break(marker)
            self.add_pending_bookmarks(marker)
        rows = tag.find_all("tr")
        if not rows:
            return
        column_count = max(sum(max(1, int(cell.get("colspan", 1))) for cell in row.find_all(["td", "th"], recursive=False)) for row in rows)
        table = self.document.add_table(rows=len(rows), cols=max(1, column_count))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        occupied: set[tuple[int, int]] = set()
        for row_index, (html_row, word_row) in enumerate(zip(rows, table.rows)):
            html_cells = html_row.find_all(["td", "th"], recursive=False)
            column = 0
            for html_cell in html_cells:
                while (row_index, column) in occupied:
                    column += 1
                colspan = max(1, int(html_cell.get("colspan", 1)))
                rowspan = max(1, int(html_cell.get("rowspan", 1)))
                end_col = min(column_count - 1, column + colspan - 1)
                end_row = min(len(rows) - 1, row_index + rowspan - 1)
                cell = table.cell(row_index, column)
                if end_col != column or end_row != row_index:
                    cell = cell.merge(table.cell(end_row, end_col))
                cell.text = ""
                paragraph = cell.paragraphs[0]
                for child in html_cell.children:
                    self.add_inline(paragraph, child, bold=html_cell.name == "th")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                fill = color_hex(self.resolver.resolve(html_cell, "background-color"))
                if fill:
                    set_cell_shading(cell, fill)
                for rr in range(row_index, end_row + 1):
                    for cc in range(column, end_col + 1):
                        occupied.add((rr, cc))
                column = end_col + 1
            if html_row.find("th"):
                set_repeat_table_header(word_row)
        self.state.table_count += 1

    def emit(self, node, *, list_level: int = 0, ordered: bool = False) -> None:
        if isinstance(node, NavigableString) or not isinstance(node, Tag) or _is_hidden(node):
            return
        role = str(node.get("role", "")).casefold()
        if role in {"doc-footnote", "doc-endnote"}:
            return
        name = node.name.casefold()
        classes = set(node.get("class", []))
        if (
            "source-page" in classes
            and not self.state.body_section_created
            and id(node) not in {self.toc_page_id, self.body_page_id}
        ):
            if self.front_source_pages_emitted:
                self.page_break_pending = True
            self.front_source_pages_emitted += 1
        if id(node) == self.toc_tag_id or "toc-table" in classes or "contents-list" in classes:
            self.insert_toc(toc_label(node))
            return
        if id(node) == self.body_heading_id:
            self.start_body()
        if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(name[1])
            chapter_break = level == 1 and self.state.body_section_created and id(node) != self.body_heading_id
            style = "Scripta Book Title" if id(node) == self.title_heading_id else f"Heading {level}"
            paragraph = self.add_paragraph(node, style)
            if chapter_break:
                paragraph.paragraph_format.page_break_before = True
            self.state.heading_count += 1
            return
        if name == "p":
            style = "Scripta Book Subtitle" if self.title_page and node.find_parent("section", class_="source-page") is self.title_page else None
            self.add_paragraph(node, style)
            return
        if name == "blockquote":
            direct = node.find_all("p", recursive=False)
            if direct:
                for child in direct:
                    self.add_paragraph(child, "Book Quote")
            else:
                self.add_paragraph(node, "Book Quote")
            return
        if name == "pre":
            paragraph = self.document.add_paragraph(style="Code Block")
            self.add_pending_bookmarks(paragraph)
            for child in node.children:
                self.add_inline(paragraph, child, code=True)
            return
        if name == "figure":
            image = node.find("img")
            if image:
                self.add_image(image)
            caption = node.find("figcaption")
            if caption:
                self.add_paragraph(caption, "Caption")
            return
        if name == "img":
            self.add_image(node)
            return
        if name == "table":
            self.add_table(node)
            return
        if name in {"ul", "ol"}:
            for item in node.find_all("li", recursive=False):
                paragraph = self.document.add_paragraph(style=("List Number" if name == "ol" else "List Bullet") + (f" {min(3, list_level + 1)}" if list_level else ""))
                for child in item.children:
                    if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                        continue
                    self.add_inline(paragraph, child)
                for child in item.find_all(["ul", "ol"], recursive=False):
                    self.emit(child, list_level=list_level + 1, ordered=child.name == "ol")
            return
        if name == "dt":
            paragraph = self.add_paragraph(node)
            for run in paragraph.runs:
                run.bold = True
            return
        if name == "dd":
            paragraph = self.add_paragraph(node)
            paragraph.paragraph_format.left_indent = Inches(0.3)
            return
        if name == "hr":
            paragraph = self.document.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            return
        direct_blocks = [child for child in node.children if isinstance(child, Tag) and child.name in BLOCKS]
        if name not in {"main", "body", "html"} and not direct_blocks and node.get_text(" ", strip=True):
            self.add_paragraph(node)
            return
        if node.get("id"):
            self.pending_bookmarks.append(str(node["id"]))
        for child in node.children:
            if isinstance(child, Tag):
                self.emit(child, list_level=list_level, ordered=ordered)

    def build(self) -> Document:
        root = self.soup.select_one("main[data-reader-content]")
        if root is None:
            raise Html2DocError("Input is not Scripta HTML: main[data-reader-content] is missing.")
        if self.body_heading_id is None and self.toc_tag_id is None:
            self.insert_toc()
        self.emit(root)
        if not self.state.toc_inserted:
            self.insert_toc()
        if not self.state.body_section_created:
            self.state.warnings.append(Finding("warning", "body-boundary-ambiguous", "No body chapter boundary was detected; front-matter page numbering remains in use."))
        return self.document


def _default_output(source: Path) -> Path:
    name = source.parent.name if source.name.casefold() == "index.html" else source.stem
    return source.with_name(f"{name}.docx")


def is_owned_docx(path: Path) -> bool:
    if not path.is_file():
        return False
    try:
        document = Document(path)
        return "html2doc-skill" in (document.core_properties.keywords or "")
    except Exception:
        return False


def convert_html(source: Path, output: Path | None = None, *, title: str | None = None, author: str | None = None, language: str | None = None, overwrite: bool = False) -> dict[str, object]:
    source = source.resolve()
    if not source.is_file():
        raise Html2DocError(f"Input HTML does not exist: {source}")
    target = (output.resolve() if output else _default_output(source).resolve())
    if target.suffix.casefold() != ".docx":
        raise Html2DocError("Output must use the .docx extension.")
    if target.exists() and (not overwrite or not is_owned_docx(target)):
        reason = "Use --overwrite for a skill-owned output." if is_owned_docx(target) else "Refusing to overwrite a DOCX not owned by html2doc-skill."
        raise Html2DocError(f"Output already exists: {target}. {reason}")
    soup = BeautifulSoup(source.read_text(encoding="utf-8"), "html.parser")
    if soup.select_one("main[data-reader-content]") is None:
        raise Html2DocError("Input is not Scripta HTML: main[data-reader-content] is missing.")
    resolver = StyleResolver(soup, source)
    metadata = _metadata(soup, source, title=title, author=author, language=language, resolver=resolver)
    builder = Builder(source, soup, resolver, metadata)
    document = builder.build()
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".docx", delete=False) as handle:
        candidate = Path(handle.name)
    try:
        document.save(candidate)
        patch_notes(candidate, builder.state.footnotes, builder.state.endnotes)
        from .validation import validate_docx
        report = validate_docx(source, candidate)
        if report["status"] == "failed":
            raise Html2DocError("Generated DOCX failed validation: " + "; ".join(item["message"] for item in report["findings"] if item["severity"] == "error"))
        candidate.replace(target)
    finally:
        candidate.unlink(missing_ok=True)
    warnings = [item for item in report["findings"] if item["severity"] == "warning"] + [item.json() for item in builder.state.warnings]
    status = "passed_with_warnings" if report["status"] == "passed" and warnings else report["status"]
    return {
        "status": status, "artifact": str(target),
        "warnings": warnings,
        "metrics": report["metrics"],
    }
